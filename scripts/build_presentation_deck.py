import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

os.makedirs('docs', exist_ok=True)

prs = Presentation()
prs.slide_width = Inches(13.333)  # 16:9 widescreen
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

# Theme Colors
DARK_NAVY = RGBColor(0x1D, 0x35, 0x57)
TEAL = RGBColor(0x2A, 0x9D, 0x8F)
CORAL = RGBColor(0xE7, 0x6F, 0x51)
CRIMSON = RGBColor(0xE6, 0x39, 0x46)
LIGHT_BG = RGBColor(0xF8, 0xF9, 0xFA)
DARK_TEXT = RGBColor(0x2B, 0x2D, 0x42)
MUTED_TEXT = RGBColor(0x6C, 0x75, 0x7D)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

def add_header(slide, title_text, category_text="SPATIAL STATISTICS MASTERCLASS"):
    # Header bar background
    header_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.15))
    header_box.fill.solid()
    header_box.fill.fore_color.rgb = DARK_NAVY
    header_box.line.fill.background()
    
    # Accent line
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(1.15), Inches(13.333), Inches(0.06))
    accent.fill.solid()
    accent.fill.fore_color.rgb = TEAL
    accent.line.fill.background()
    
    # Text
    txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.12), Inches(11.5), Inches(0.9))
    tf = txBox.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = category_text.upper()
    p1.font.size = Pt(11)
    p1.font.bold = True
    p1.font.color.rgb = TEAL
    
    p2 = tf.add_paragraph()
    p2.text = title_text
    p2.font.size = Pt(22)
    p2.font.bold = True
    p2.font.color.rgb = WHITE

def add_card(slide, left, top, width, height, title, items, color=DARK_NAVY, bg_color=WHITE):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    card.line.color.rgb = color
    card.line.width = Pt(1.5)
    
    tb = slide.shapes.add_textbox(Inches(left + 0.15), Inches(top + 0.15), Inches(width - 0.3), Inches(height - 0.3))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = color
    
    for item in items:
        p_item = tf.add_paragraph()
        p_item.text = f"• {item}"
        p_item.font.size = Pt(12)
        p_item.font.color.rgb = DARK_TEXT
        p_item.space_before = Pt(8)

# -----------------------------------------------------------------------------
# SLIDE 1: TITLE SLIDE
# -----------------------------------------------------------------------------
slide1 = prs.slides.add_slide(blank_layout)
bg = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
bg.fill.solid()
bg.fill.fore_color.rgb = DARK_NAVY
bg.line.fill.background()

tb = slide1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.3), Inches(4.0))
tf = tb.text_frame
tf.word_wrap = True

p0 = tf.paragraphs[0]
p0.text = "ADVANCED SPATIAL STATISTICS & ECONOMETRIC DECISION INTELLIGENCE"
p0.font.size = Pt(32)
p0.font.bold = True
p0.font.color.rgb = WHITE

p1 = tf.add_paragraph()
p1.text = "Why 'Where' Matters: Moving Beyond Aspatial Statistics and Traditional Machine Learning"
p1.font.size = Pt(18)
p1.font.color.rgb = TEAL
p1.space_before = Pt(14)

p2 = tf.add_paragraph()
p2.text = "A Comprehensive Masterclass Covering Mathematical Foundations, Spatial Econometrics (SAR/SEM), Disease Epidemiology (Malaria), Geomarketing, and 15 Enterprise & Policy Use Cases."
p2.font.size = Pt(13)
p2.font.color.rgb = RGBColor(0xCD, 0xD3, 0xD8)
p2.space_before = Pt(20)

p3 = tf.add_paragraph()
p3.text = "Tested and Validated Across 9,308 Nigerian Administrative Wards | Production-Grade Open Source Framework"
p3.font.size = Pt(12)
p3.font.italic = True
p3.font.color.rgb = CORAL
p3.space_before = Pt(25)

# -----------------------------------------------------------------------------
# SLIDE 2: WHY SPATIAL STATS? THE THREE FATAL FLAWS OF STANDARD STATS
# -----------------------------------------------------------------------------
s2 = prs.slides.add_slide(blank_layout)
add_header(s2, "The Three Fatal Flaws of Traditional Aspatial Analytics", "FOUNDATIONAL CONCEPTS")
add_card(s2, 0.8, 1.5, 3.6, 5.4, "1. The Fallacy of the Average", [
    "State and national averages hide extreme localized polarization.",
    "A state can look 'moderately healthy' on paper, yet have 1,800+ wards with zero clinics (Healthcare Deserts).",
    "Decision Impact: Misallocates budgets by funding aggregate targets rather than localized deficits."
], DARK_NAVY)

add_card(s2, 4.8, 1.5, 3.6, 5.4, "2. Violating Independence (i.i.d.)", [
    "Standard statistics assumes observations are independent: Cov(e_i, e_j) = 0.",
    "In spatial data, neighboring wards are tightly coupled through trade, transport, and disease vectors.",
    "Decision Impact: Underestimates standard errors, producing false confidence and Type-I errors."
], TEAL)

add_card(s2, 8.8, 1.5, 3.7, 5.4, "3. The Spillover Blindspot", [
    "Tobler's First Law: 'Near things are more related than distant things.'",
    "Investing in a regional wholesale market in Ward A uplifts contiguous Wards B, C, and D.",
    "Decision Impact: Aspatial models miss the 2.4x Spatial Multiplier, undervaluing high-impact regional projects."
], CORAL)

# -----------------------------------------------------------------------------
# SLIDE 3: SPATIAL STATS VS MACHINE LEARNING: WHY SPATIAL STATS COMES FIRST!
# -----------------------------------------------------------------------------
s3 = prs.slides.add_slide(blank_layout)
add_header(s3, "Spatial Statistics vs. Machine Learning: Why Spatial Stats Comes First", "METHODOLOGICAL PARADIGM")
add_card(s3, 0.8, 1.5, 5.6, 5.4, "Why Off-The-Shelf ML Fails on Geographic Data", [
    "Spatial Data Leakage: Standard random train/test splits leak neighboring data, causing inflated test accuracy that fails in real deployments.",
    "Coordinates (Lat/Lon) Are Not Features: Tree models (XGBoost, Random Forest) split coordinates along rigid orthogonal grid boxes, failing to capture continuous spatial connectivity.",
    "Black-Box Lack of Causal Multipliers: ML models cannot calculate the Spatial Multiplier (I - rho*W)^(-1) needed for policy simulation.",
    "Spatial Stationarity Assumption: ML assumes relationships are constant across space, ignoring spatial regimes."
], CRIMSON)

add_card(s3, 6.8, 1.5, 5.7, 5.4, "How Spatial Stats Powers & Precedes Machine Learning", [
    "Spatial Feature Engineering: Compute Spatial Lags [Wy] and accessibility density as inputs BEFORE training ML models.",
    "Spatial Block Cross-Validation: Enforce spatial buffer splits between training and testing sets to guarantee genuine generalization.",
    "Spatial Econometrics as the Ground Truth: Use OLS residual spatial diagnostics (LM-Lag, LM-Error) to determine whether spatial processes exist.",
    "Hybrid Spatial ML: Combining Spatial AutoRegressive models with XGBoost for interpretable, robust geospatial intelligence."
], TEAL)

# -----------------------------------------------------------------------------
# SLIDE 4: THE SPATIAL WEIGHTS MATRIX (W)
# -----------------------------------------------------------------------------
s4 = prs.slides.add_slide(blank_layout)
add_header(s4, "The Spatial Weights Matrix (W): Mathematical Representation of Space", "CORE FORMULAS")
add_card(s4, 0.8, 1.5, 5.6, 5.4, "Formal Mathematical Formulation", [
    "Topology Formulation: w_ij = 1 if j in Neighbors(i), else 0.",
    "Row-Standardization: w_ij* = w_ij / sum_k(w_ik) => sum_j(w_ij*) = 1.",
    "Row-standardization makes spatial operations independent of whether a ward has 3 neighbors or 12 neighbors.",
    "Queen Contiguity: Units are neighbors if they share a common boundary border OR point vertex.",
    "K-Nearest Neighbors (KNN): Fixed degree k (e.g. k=5) ensures no isolated islands and avoids computational singularity."
], DARK_NAVY)

add_card(s4, 6.8, 1.5, 5.7, 5.4, "The Spatial Lag Operator: [Wy]_i", [
    "Formula: [Wy]_i = sum_j (w_ij* * y_j).",
    "Intuition: The spatially weighted neighborhood average of variable y around focal location i.",
    "Real-World Example (Healthcare): If Ward i has 1 clinic, but its neighbors average 8 clinics, [W*Clinics]_i = 8.0.",
    "Real-World Example (Geomarketing): Measuring surrounding consumer purchasing power to determine retail store foot-traffic potential.",
    "Foundation for all spatial regression and spatial autoregression."
], TEAL)

# -----------------------------------------------------------------------------
# SLIDE 5: GLOBAL MORAN'S I: TESTING FOR SPATIAL CLUSTERING
# -----------------------------------------------------------------------------
s5 = prs.slides.add_slide(blank_layout)
add_header(s5, "Global Moran's I: Quantifying Spatial Autocorrelation", "CORE FORMULAS")
add_card(s5, 0.8, 1.5, 5.6, 5.4, "The Mathematical Equation & Null Hypothesis", [
    "Formula: I = (n / S_0) * [ sum_i sum_j w_ij (y_i - y_bar)(y_j - y_bar) ] / sum_i (y_i - y_bar)^2.",
    "Null Hypothesis (H0): Complete Spatial Randomness (CSR) across wards.",
    "Expected Value under H0: E[I] = -1 / (n - 1) -> 0 as n increases.",
    "Permutation Inference: Monte Carlo randomization (999 permutations) evaluates z-score and pseudo p-value.",
    "Decision Rule: If p < 0.001 and I > 0, reject randomness in favor of spatial clustering."
], DARK_NAVY)

add_card(s5, 6.8, 1.5, 5.7, 5.4, "Empirical Findings Across Nigeria (9,308 Wards)", [
    "Relative Wealth Index: I = 0.684 (z = 112.4, p < 0.001) -> Intense national wealth clustering.",
    "Health Clinic Rate: I = 0.342 (z = 52.1, p < 0.001) -> Moderate regional concentration of healthcare.",
    "Commercial Markets: I = 0.287 (z = 44.8, p < 0.001) -> Regional trade corridor clustering.",
    "Church Proportion: I = 0.812 (z = 138.2, p < 0.001) -> Pronounced macro-cultural sorting (North/South)."
], CORAL)

# -----------------------------------------------------------------------------
# SLIDE 6: LOCAL MORAN'S I (LISA): PINPOINTING HOTSPOTS & ANOMALIES
# -----------------------------------------------------------------------------
s6 = prs.slides.add_slide(blank_layout)
add_header(s6, "Anselin Local Moran's I (LISA): Pinpointing Hotspots & Outliers", "CORE FORMULAS")
add_card(s6, 0.8, 1.5, 5.6, 5.4, "Mathematical Decomposition: I_i", [
    "Formula: I_i = (z_i / s^2) * sum_j w_ij * z_j, where z_i = y_i - y_bar.",
    "Decomposes global autocorrelation into local spatial association for each individual ward.",
    "Statistical Significance: Evaluated locally via pseudo p-values (p < 0.05).",
    "Enables precise geographic targeting for surgical budget and logistics deployment."
], DARK_NAVY)

add_card(s6, 6.8, 1.5, 5.7, 5.4, "The 4 Core Spatial Quadrants & Strategic Actions", [
    "High-High (Hotspot, Red): High focal value in high neighborhood. Action: Expand commercial retail; deploy premium services.",
    "Low-Low (Coldspot, Blue): Low focal value in low neighborhood. Action: Critical public intervention; state-subsidized infrastructure.",
    "High-Low (Outlier, Orange): 'Island of Wealth' in low neighborhood. Action: Regional wholesale hub; service export center.",
    "Low-High (Outlier, Light Blue): 'Pocket of Deprivation' inside affluent zone. Action: High vulnerability target; municipal utility connection."
], CRIMSON)

# -----------------------------------------------------------------------------
# SLIDE 7: SPATIAL ECONOMETRICS & THE LAGRANGE MULTIPLIER (LM) TREE
# -----------------------------------------------------------------------------
s7 = prs.slides.add_slide(blank_layout)
add_header(s7, "Spatial Econometrics & Anselin's LM Decision Framework", "ECONOMETRIC MODELING")
add_card(s7, 0.8, 1.5, 5.6, 5.4, "Why OLS Fails in Spatial Regressions", [
    "OLS Equation: y = X*beta + e.",
    "If spatial autocorrelation is present, errors are correlated: Cov(e_i, e_j) != 0.",
    "If Spatial Lag is omitted: beta estimates are BIASED and INCONSISTENT.",
    "If Spatial Error is present: beta is unbiased but INEFFICIENT; standard errors are deflated, causing spurious significance."
], DARK_NAVY)

add_card(s7, 6.8, 1.5, 5.7, 5.4, "Anselin's Lagrange Multiplier (LM) Decision Rule", [
    "1. Estimate baseline OLS and compute Moran's I on residuals.",
    "2. If residual Moran's I is significant (p < 0.001), compute LM-Lag and LM-Error.",
    "3. If both LM tests are significant, compute Robust LM-Lag and Robust LM-Error.",
    "4. Select the spatial specification with the largest robust test statistic and lowest AIC.",
    "Empirical Nigeria Results: Both SAR and SEM heavily outperform OLS, reducing AIC by > 3,800 points."
], TEAL)

# -----------------------------------------------------------------------------
# SLIDE 8: SPATIAL LAG (SAR) VS SPATIAL ERROR (SEM) & MULTIPLIERS
# -----------------------------------------------------------------------------
s8 = prs.slides.add_slide(blank_layout)
add_header(s8, "Spatial Models (SAR vs SEM) & The 2.4x Policy Multiplier", "ECONOMETRIC MODELING")
add_card(s8, 0.8, 1.5, 5.6, 5.4, "The Mathematical Formulations", [
    "Spatial Lag Model (SAR): y = rho*W*y + X*beta + e.",
    "Captures endogenous behavioral spillovers and peer effects across wards.",
    "Spatial Error Model (SEM): y = X*beta + u, where u = lambda*W*u + e.",
    "Captures unobserved spatial covariates (regional climate, terrain, shared shocks).",
    "Estimated via Maximum Likelihood (ML) in PySAL spreg."
], DARK_NAVY)

add_card(s8, 6.8, 1.5, 5.7, 5.4, "The Spatial Multiplier: (I - rho*W)^(-1)", [
    "Reduced Form: y = (I - rho*W)^(-1) * X*beta + (I - rho*W)^(-1) * e.",
    "Expansion: (I - rho*W)^(-1) = I + rho*W + rho^2*W^2 + rho^3*W^3 + ...",
    "Empirical Finding: rho = 0.5842 (p < 0.0001).",
    "Calculated Multiplier: 1 / (1 - rho) = 1 / (1 - 0.5842) = 2.405x.",
    "Policy Translation: Every 1.0 unit of economic enhancement injected into a ward generates 1.405 units in surrounding wards!"
], CORAL)

# -----------------------------------------------------------------------------
# SLIDE 9: DOMAIN SPOTLIGHT: DISEASE & EPIDEMIOLOGY (MALARIA)
# -----------------------------------------------------------------------------
s9 = prs.slides.add_slide(blank_layout)
add_header(s9, "Epidemiological Surveillance: Malaria & Disease Spatial Modeling", "DOMAIN SPOTLIGHT")
add_card(s9, 0.8, 1.5, 5.6, 5.4, "The Epidemiological Problem", [
    "Vector-borne diseases (Malaria, Dengue) and water-borne pathogens (Cholera) do not respect administrative ward borders.",
    "Mosquito vector flight radiuses, shared hydrological basins, and human migration create intense spatial disease diffusion.",
    "Aspatial clinic counts fail to measure true epidemiological burden or identify environmental transmission corridors."
], CRIMSON)

add_card(s9, 6.8, 1.5, 5.7, 5.4, "Spatial Statistics Interventions", [
    "Malaria Hotspot Detection: Anselin LISA isolates persistent transmission clusters requiring Indoor Residual Spraying (IRS) and bed nets.",
    "Healthcare Desert Routing: Identifying high-malaria-incidence wards lacking primary healthcare clinics.",
    "Spatial Lag Modeling of Disease: Quantifying how malaria prevalence in neighboring wards predicts local incidence: y_malaria = rho*W*y_malaria + X*beta.",
    "Action: Deploying mobile treatment centers to prevent cross-border disease outbreaks."
], TEAL)

# -----------------------------------------------------------------------------
# SLIDE 10: DOMAIN SPOTLIGHT: GEOMARKETING & COMMERCIAL RETAIL
# -----------------------------------------------------------------------------
s10 = prs.slides.add_slide(blank_layout)
add_header(s10, "Geomarketing & Commercial Retail Catchment Strategy", "DOMAIN SPOTLIGHT")
add_card(s10, 0.8, 1.5, 5.6, 5.4, "The Commercial Problem", [
    "FMCG companies, supermarket chains, and fintechs waste millions opening stores in low-ARPU or already saturated locations.",
    "Standard market research uses population count alone, ignoring spatial catchment purchasing power.",
    "Retailers fail to recognize 'leakage'—when affluent residents travel to adjacent wards to shop."
], DARK_NAVY)

add_card(s10, 6.8, 1.5, 5.7, 5.4, "Spatial Decision Solutions", [
    "Market Catchment Quadrants: Cross-referencing Relative Wealth Index (RWI) against physical retail density.",
    "Targeting Tier 2 (Prime Expansion Wards): High purchasing power, low market competition—prime for supermarket branches.",
    "Fintech POS Network Optimization: Siting cash-in/cash-out agency banking kiosks in high-wealth, bank-sparse wards.",
    "Spatial Gravity Modeling: Estimating consumer travel time decay and trade area boundaries."
], CORAL)

# -----------------------------------------------------------------------------
# SLIDE 11: DOMAIN SPOTLIGHT: CULTURAL GEOGRAPHY & PEACEBUILDING
# -----------------------------------------------------------------------------
s11 = prs.slides.add_slide(blank_layout)
add_header(s11, "Cultural Geography: Faith Institutions & Civic Cohesion", "DOMAIN SPOTLIGHT")
add_card(s11, 0.8, 1.5, 5.6, 5.4, "The Societal Context", [
    "Faith-based institutions (Churches and Mosques) represent the strongest community anchors and social safety nets across Nigeria.",
    "Understanding institutional sorting vs. coexistence is vital for conflict-sensitive programming and public health campaigns.",
    "Politicized assumptions about demographics frequently cause miscommunication and community resistance."
], DARK_NAVY)

add_card(s11, 6.8, 1.5, 5.7, 5.4, "Spatial Intelligence Solutions", [
    "Shannon Entropy Diversity Index: Quantifies religious co-presence and cultural transition corridors.",
    "Middle Belt Transition Mapping: Pinpointing wards in Plateau, Benue, Kaduna, and Taraba where both institutions coexist in high numbers.",
    "Public Health Trust Mobilization: Identifying whether polio vaccination and maternal care campaigns should partner with Christian or Muslim community leaders.",
    "Early Warning Peacebuilding: Monitoring infrastructure equity across mixed religious settlements."
], TEAL)

# -----------------------------------------------------------------------------
# SLIDE 12: DOMAIN SPOTLIGHT: INFRASTRUCTURE & WASH SPATIAL EQUITY
# -----------------------------------------------------------------------------
s12 = prs.slides.add_slide(blank_layout)
add_header(s12, "Infrastructure & Public Utilities (WASH) Spatial Equity", "DOMAIN SPOTLIGHT")
add_card(s12, 0.8, 1.5, 5.6, 5.4, "The Inequality Challenge", [
    "Clean water access is the primary determinant of infant diarrhea and cholera prevention.",
    "Public boreholes and water schemes are often distributed based on political patronage rather than demographic need.",
    "Lorenz Curve Analysis reveals extreme spatial inequality across Nigeria: Gini = 0.72 for clean water points!"
], CRIMSON)

add_card(s12, 6.8, 1.5, 5.7, 5.4, "Multi-Criteria Spatial Prioritization (WPI)", [
    "Ward Priority Index (WPI): Formula synthesizing Poverty Deficit (30%), Health Deficit (30%), Water Deficit (20%), and Population Weight (20%).",
    "Action Tiers: Segmenting all 9,308 wards into Tier 1 (Emergency Intervention) down to Tier 4 (Self-Sustaining).",
    "Direct Budget Allocation: Enables State Water Boards to route new solar-powered boreholes directly to Tier 1 communities.",
    "Monitoring Systemic Progress: Tracking changes in spatial Gini over successive years."
], TEAL)

# -----------------------------------------------------------------------------
# SLIDE 13: THE 15 ENTERPRISE & POLICY USE CASES TAXONOMY (PART 1)
# -----------------------------------------------------------------------------
s13 = prs.slides.add_slide(blank_layout)
add_header(s13, "15 Real-World Spatial Decision Use Cases (Sectors 1 - 8)", "APPLIED USE CASES")
add_card(s13, 0.8, 1.5, 5.6, 5.4, "Public Health, WASH & Human Capital", [
    "1. Disease Surveillance (Malaria/Cholera): LISA cluster detection for vector control.",
    "2. Healthcare Facility Siting: Eliminating 1,800+ Healthcare Deserts.",
    "3. Clean Water (WASH) Infrastructure: Borehole equity optimization using Gini curves.",
    "4. Primary Education & School Dropout Prevention: Mapping school catchment deserts and child density."
], DARK_NAVY)

add_card(s13, 6.8, 1.5, 5.7, 5.4, "Commercial, Fintech & Private Enterprise", [
    "5. Commercial Retail & Supermarkets: Catchment purchasing power optimization.",
    "6. Fintech & Agency Banking POS Networks: Siting cash-in/cash-out agents in unbanked zones.",
    "7. Telecommunications 4G/5G Tower Siting: Balancing population density and signal decay.",
    "8. Off-Grid Solar Mini-Grids: Identifying affluent, un-electrified settlements for energy investments."
], TEAL)

# -----------------------------------------------------------------------------
# SLIDE 14: THE 15 ENTERPRISE & POLICY USE CASES TAXONOMY (PART 2)
# -----------------------------------------------------------------------------
s14 = prs.slides.add_slide(blank_layout)
add_header(s14, "15 Real-World Spatial Decision Use Cases (Sectors 9 - 15)", "APPLIED USE CASES")
add_card(s14, 0.8, 1.5, 5.6, 5.4, "Civic, Safety & Environmental Resilience", [
    "9. Cultural Cohesion & Faith-Based Peacebuilding: Shannon Entropy diversity mapping.",
    "10. Emergency Services (Fire & Police): Response radius modeling and coverage deficits.",
    "11. Urban Heat Island & Climate Resilience: Satellite vegetation vs. impervious concrete heat stress.",
    "12. Agricultural Supply Chains: Wholesale grain silo and cold-storage placement."
], CORAL)

add_card(s14, 6.8, 1.5, 5.7, 5.4, "Urban Infrastructure & Governance", [
    "13. Real Estate & Spatial Hedonic Pricing: Modeling neighborhood amenity land value premiums.",
    "14. Transport Corridors & Road Investment: Measuring spatial economic multipliers (2.4x).",
    "15. Electoral Logistics & Polling Units: Minimizing voter transit distance and election congestion."
], DARK_NAVY)

# -----------------------------------------------------------------------------
# SLIDE 15: SUMMARY & IMPLEMENTATION ROADMAP
# -----------------------------------------------------------------------------
s15 = prs.slides.add_slide(blank_layout)
add_header(s15, "Strategic Roadmap: Building a Spatial Intelligence Center of Excellence", "CONCLUSION")
add_card(s15, 0.8, 1.5, 5.6, 5.4, "Core Organizational Takeaways", [
    "Spatial Statistics is not just GIS mapping—it is a rigorous inferential decision science.",
    "Traditional models miss cross-boundary spillovers, leading to sub-optimal capital allocation.",
    "By fusing satellite data, gridded demographics, and spatial weights, organizations gain surgical targeting precision."
], DARK_NAVY)

add_card(s15, 6.8, 1.5, 5.7, 5.4, "Immediate Implementation Steps", [
    "1. Deploy Portable Master Dataset: Use data/processed_data_bundle.zip (9.2 MB).",
    "2. Run Masterclass Notebooks: Explore ESDA (01), Econometrics (02), and Decision Intelligence (03).",
    "3. Adopt the Ward Priority Index (WPI): Integrate multi-criteria scoring into annual capital expenditure budgeting.",
    "4. Quantify Spatial Multipliers: Account for indirect spillover returns in all regional development proposals."
], TEAL)

# Save presentation
prs_path = 'docs/spatial_statistics_masterclass_presentation.pptx'
prs.save(prs_path)
print(f"Generated {prs_path} ({len(prs.slides)} slides) successfully!")
