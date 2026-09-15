import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

os.makedirs('docs', exist_ok=True)

prs = Presentation()
prs.slide_width = Inches(13.333)  # 16:9 widescreen
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

# Theme Colors
NAVY = RGBColor(0x1D, 0x35, 0x57)
TEAL = RGBColor(0x2A, 0x9D, 0x8F)
CORAL = RGBColor(0xE7, 0x6F, 0x51)
CRIMSON = RGBColor(0xD9, 0x04, 0x29)
LIGHT_BG = RGBColor(0xF8, 0xF9, 0xFA)
DARK_TEXT = RGBColor(0x21, 0x25, 0x29)
MUTED = RGBColor(0x6C, 0x75, 0x7D)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BOX_BG = RGBColor(0xEE, 0xF4, 0xF8)

def add_header(slide, title, category="SPATIAL STATISTICS"):
    # Header bar
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.1))
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    
    # Accent line
    acc = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(1.1), Inches(13.333), Inches(0.05))
    acc.fill.solid()
    acc.fill.fore_color.rgb = TEAL
    acc.line.fill.background()
    
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.12), Inches(11.7), Inches(0.9))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p1 = tf.paragraphs[0]
    p1.text = category.upper()
    p1.font.size = Pt(11)
    p1.font.bold = True
    p1.font.color.rgb = TEAL
    
    p2 = tf.add_paragraph()
    p2.text = title
    p2.font.size = Pt(21)
    p2.font.bold = True
    p2.font.color.rgb = WHITE

def add_text_box(slide, left, top, width, height, title, lines, color=NAVY, bg=WHITE):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    card.fill.solid()
    card.fill.fore_color.rgb = bg
    card.line.color.rgb = color
    card.line.width = Pt(1.5)
    
    tb = slide.shapes.add_textbox(Inches(left + 0.2), Inches(top + 0.2), Inches(width - 0.4), Inches(height - 0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = color
    
    for l in lines:
        p_l = tf.add_paragraph()
        p_l.text = l
        p_l.font.size = Pt(12)
        p_l.font.color.rgb = DARK_TEXT
        p_l.space_before = Pt(8)

def add_image_safe(slide, img_path, left, top, width, height):
    if os.path.exists(img_path):
        slide.shapes.add_picture(img_path, Inches(left), Inches(top), Inches(width), Inches(height))

# =============================================================================
# SLIDE 1: TITLE
# =============================================================================
s1 = prs.slides.add_slide(blank_layout)
bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
bg1.fill.solid()
bg1.fill.fore_color.rgb = NAVY
bg1.line.fill.background()

tb = s1.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.3), Inches(4.0))
tf = tb.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "ADVANCED SPATIAL STATISTICS & DECISION SCIENCE"
p.font.size = Pt(34)
p.font.bold = True
p.font.color.rgb = WHITE

p2 = tf.add_paragraph()
p2.text = "How Location Data and Spatial Modeling Drive Smarter Decisions Across Every Sector"
p2.font.size = Pt(18)
p2.font.color.rgb = TEAL
p2.space_before = Pt(14)

p3 = tf.add_paragraph()
p3.text = "Every sector starts with a question. Spatial statistics gives the answer.\nCovering Disease Epidemiology (Malaria), Health Deserts, Geomarketing, Cultural Geography, and Spatial Econometrics across 9,308 Nigerian Wards."
p3.font.size = Pt(13)
p3.font.color.rgb = RGBColor(0xCD, 0xD3, 0xD8)
p3.space_before = Pt(20)

# =============================================================================
# SLIDE 2: THE BIG IDEA - WHY "WHERE" MATTERS
# =============================================================================
s2 = prs.slides.add_slide(blank_layout)
add_header(s2, "The Big Idea: Why Traditional Statistics Fails in the Real World", "CORE FOUNDATIONS")
add_text_box(s2, 0.8, 1.4, 5.7, 5.6, "The Fallacy of the Average", [
    "Traditional analytics uses state or national averages.",
    "A state can look moderately healthy on paper, while having hundreds of wards with ZERO clinics.",
    "Decisions based on averages waste money because they fund the wrong places.",
    "Location is not just another column in a spreadsheet—it dictates who has access and who is left behind."
], CRIMSON)
add_text_box(s2, 6.8, 1.4, 5.7, 5.6, "Tobler's First Law of Geography", [
    "'Everything is related to everything else, but near things are more related than distant things.' — Waldo Tobler (1970)",
    "Events in one ward do not stay in that ward.",
    "Malaria mosquitoes fly across borders.",
    "Shoppers travel to neighboring markets.",
    "When you invest in Ward A, surrounding Wards B, C, and D also benefit (Spatial Spillovers)."
], TEAL)

# =============================================================================
# SLIDE 2B: SPATIAL STATS VS TRADITIONAL STATS (HEAD-TO-HEAD COMPARISON)
# =============================================================================
s2b = prs.slides.add_slide(blank_layout)
add_header(s2b, "Spatial Statistics vs. Traditional Statistics: Head-to-Head Comparison", "CORE COMPARISON")
add_text_box(s2b, 0.8, 1.4, 5.7, 5.6, "Traditional Aspatial Statistics", [
    "Core Assumption: Observations are independent and identically distributed (i.i.d.).",
    "Correlation: Pearson/Spearman assumes pairs are isolated in space.",
    "Regression: Standard OLS y = X*beta + e assumes zero error correlation Cov(e_i, e_j) = 0.",
    "Outliers: Measured purely by global z-score (> 3 standard deviations from national mean).",
    "Multiplier: 1.0x (Assumes an investment only impacts the targeted boundary).",
    "Limitation: Blind to geographic clustering, spillovers, and boundary effects."
], CRIMSON)
add_text_box(s2b, 6.8, 1.4, 5.7, 5.6, "Advanced Spatial Statistics", [
    "Core Foundation: Embraces Tobler's Law—neighboring units are systematically coupled.",
    "Correlation: Global & Bivariate Moran's I accounts for geographic distance and contiguity.",
    "Regression: Spatial Lag (SAR) and Spatial Error (SEM) capture spatial feedback loops.",
    "Outliers: Anselin LISA detects localized anomalies (High-Low and Low-High pockets).",
    "Multiplier: 2.4x (Explicitly measures how investments spill into contiguous wards).",
    "Advantage: Eliminates biased standard errors and unlocks true surgical precision."
], TEAL)

# =============================================================================
# SLIDE 2C: THE 4 MAJOR ADVANTAGES OF SPATIAL STATISTICS
# =============================================================================
s2c = prs.slides.add_slide(blank_layout)
add_header(s2c, "The 4 Strategic Advantages of Spatial Statistics in Decision-Making", "CORE ADVANTAGES")
add_text_box(s2c, 0.8, 1.4, 2.75, 5.6, "1. Bias Elimination", [
    "Prevents Omitted Variable Bias.",
    "In traditional regressions, spatial spillovers hide in the error term, distorting beta coefficients.",
    "Spatial models isolate genuine policy effects from geographic artifacts."
], NAVY)
add_text_box(s2c, 3.8, 1.4, 2.75, 5.6, "2. True Significance", [
    "Stops Type-I Errors.",
    "Non-spatial models produce deflated standard errors and artificially inflated t-statistics.",
    "Spatial modeling proves whether a policy effect is real or just shared regional noise."
], TEAL)
add_text_box(s2c, 6.8, 1.4, 2.75, 5.6, "3. Surgical Targeting", [
    "Bypasses Crude Quotas.",
    "Instead of dividing money equally across 774 LGAs, Anselin LISA isolates the exact 1,800+ desert wards that need emergency clinics.",
    "Zero wasted capital."
], CORAL)
add_text_box(s2c, 9.8, 1.4, 2.75, 5.6, "4. Quantified Spillovers", [
    "Unlocks the 2.4x Multiplier.",
    "Shows ministers and CEOs that spending $1M in Ward A creates $1.4M of indirect economic activity in contiguous wards.",
    "Justifies regional co-investments."
], CRIMSON)
s3 = prs.slides.add_slide(blank_layout)
add_header(s3, "Why Spatial Statistics Must Come BEFORE Machine Learning", "METHODOLOGY")
add_text_box(s3, 0.8, 1.4, 5.7, 5.6, "The Trap: Spatial Data Leakage in ML", [
    "If you train XGBoost or Random Forest on spatial data with random train/test splits, the model cheats.",
    "Test points neighbor training points, creating artificial 95% accuracy that completely fails when deployed in a new region.",
    "Tree models cannot understand continuous distance—they chop coordinates into rigid square boxes.",
    "Off-the-shelf ML cannot give you causal policy multipliers."
], CRIMSON)
add_text_box(s3, 6.8, 1.4, 5.7, 5.6, "The Solution: Spatial Stats First", [
    "1. Spatial Feature Engineering: Calculate Spatial Lags [Wy] and distance decay before feeding data to ML.",
    "2. Spatial Block Cross-Validation: Split data by geographic clusters to prevent data leakage.",
    "3. Spatial Econometrics: Test whether spatial spillovers exist using Lagrange Multiplier (LM) tests.",
    "Spatial statistics provides the foundation that makes Machine Learning valid and reliable."
], TEAL)

# =============================================================================
# SLIDE 4: END-TO-END PIPELINE FLOWCHART (WITH IMAGE)
# =============================================================================
s4 = prs.slides.add_slide(blank_layout)
add_header(s4, "The End-to-End Decision Pipeline: From Raw Data to Actions", "PIPELINE ARCHITECTURE")
add_image_safe(s4, 'docs/figures/00_architecture_and_methodology_flowchart.png', 0.8, 1.4, 11.7, 5.6)

# =============================================================================
# SLIDE 5: FORMULA 1 - SPATIAL WEIGHTS MATRIX (W)
# =============================================================================
s5 = prs.slides.add_slide(blank_layout)
add_header(s5, "Formula 1: The Spatial Weights Matrix (W)", "MATHEMATICAL TOOLS")
add_text_box(s5, 0.8, 1.4, 5.7, 5.6, "How We Connect Geography Mathematically", [
    "How do we tell a computer which wards are neighbors?",
    "We create a matrix W where row i and column j represent two wards.",
    "Unstandardized Adjacency:",
    "   w_ij = 1  (if Ward i and Ward j share a boundary)",
    "   w_ij = 0  (otherwise)",
    "Row-Standardization:",
    "   w*_ij = w_ij / sum_k(w_ik)",
    "This makes the weights sum to 1.0 for every ward, so a ward with 10 neighbors is on the same scale as a ward with 3 neighbors."
], DARK_TEXT)
add_text_box(s5, 6.8, 1.4, 5.7, 5.6, "Queen vs. K-Nearest Neighbors (KNN)", [
    "Queen Contiguity:",
    "Wards are neighbors if they share an edge or a corner point.",
    "K-Nearest Neighbors (KNN-5):",
    "Connects each ward to its 5 closest neighbors.",
    "Why KNN is essential in practice:",
    "It prevents isolated islands from breaking the math, ensuring every single ward has valid neighbors for modeling."
], TEAL)

# =============================================================================
# SLIDE 6: FORMULA 2 - THE SPATIAL LAG OPERATOR
# =============================================================================
s6 = prs.slides.add_slide(blank_layout)
add_header(s6, "Formula 2: The Spatial Lag Operator [Wy]", "MATHEMATICAL TOOLS")
add_text_box(s6, 0.8, 1.4, 5.7, 5.6, "The Formula & Definition", [
    "The Spatial Lag of variable y for Ward i is:",
    "   [Wy]_i = sum_j ( w*_ij * y_j )",
    "In simple words:",
    "[Wy]_i is the average value of y in the neighborhood around Ward i.",
    "If y is the number of clinics, [Wy]_i tells you whether surrounding wards have many clinics or zero clinics."
], NAVY)
add_text_box(s6, 6.8, 1.4, 5.7, 5.6, "Real-World Examples", [
    "In Public Health:",
    "Ward i has 0 clinics, but [W*Clinics] = 6.0. Patients can walk to neighboring wards.",
    "Ward i has 0 clinics and [W*Clinics] = 0.0. This is a severe, isolated Healthcare Desert.",
    "In Retail Marketing:",
    "Measuring the average wealth in surrounding wards [W*Wealth] tells a store whether the broader area has strong purchasing power."
], TEAL)

# =============================================================================
# SLIDE 7: FORMULA 3 - GLOBAL MORAN'S I (TESTING CLUSTERING)
# =============================================================================
s7 = prs.slides.add_slide(blank_layout)
add_header(s7, "Formula 3: Global Moran's I (Testing for Spatial Clustering)", "MATHEMATICAL TOOLS")
add_text_box(s7, 0.8, 1.4, 5.7, 5.6, "The Mathematical Formula", [
    "I = (n / S_0) * [ sum_i sum_j w_ij (y_i - y_bar)(y_j - y_bar) ] / sum_i (y_i - y_bar)^2",
    "Null Hypothesis (H0):",
    "Values are scattered randomly across the country.",
    "Expected Value if Random:",
    "   E[I] = -1 / (n - 1)  (approx 0)",
    "Rules of Interpretation:",
    "• I > 0 and p < 0.05: Strong Clustering (High near High, Low near Low).",
    "• I < 0 and p < 0.05: Dispersion (Checkerboard pattern).",
    "• p >= 0.05: Pure Random Noise."
], DARK_TEXT)
add_text_box(s7, 6.8, 1.4, 5.7, 5.6, "Results Across 9,308 Nigerian Wards", [
    "Wealth Index (RWI):",
    "I = 0.684, z = 112.4, p = 0.001 -> Extreme wealth clustering.",
    "Malaria Prevalence:",
    "I = 0.886, z = 140.5, p = 0.001 -> Intense geographic transmission corridors.",
    "Religious Institutions (Churches):",
    "I = 0.812, z = 138.2, p = 0.001 -> Severe North-South cultural divide.",
    "Conclusion: Random models are completely wrong for Nigeria."
], CORAL)

# =============================================================================
# SLIDE 8: FORMULA 4 - LOCAL MORAN'S I (LISA CLUSTERS)
# =============================================================================
s8 = prs.slides.add_slide(blank_layout)
add_header(s8, "Formula 4: Local Moran's I (LISA) & Hotspot Detection", "MATHEMATICAL TOOLS")
add_text_box(s8, 0.8, 1.4, 5.7, 5.6, "The Local Equation", [
    "Global Moran tells us clustering exists. Local Moran tells us WHERE it is.",
    "Formula for Ward i:",
    "   I_i = (z_i / s^2) * sum_j ( w_ij * z_j )",
    "where z_i = y_i - y_bar (deviation from national mean).",
    "Each ward gets its own score and statistical test.",
    "Wards are assigned into one of 4 quadrants based on their own value and their neighbors' values."
], NAVY)
add_text_box(s8, 6.8, 1.4, 5.7, 5.6, "The 4 Core Quadrants", [
    "High-High (Hotspot - Red):",
    "High value surrounded by high neighbors.",
    "Low-Low (Coldspot - Blue):",
    "Low value surrounded by low neighbors.",
    "High-Low (Outlier - Orange):",
    "High value surrounded by low neighbors ('Island of Wealth').",
    "Low-High (Outlier - Light Blue):",
    "Low value surrounded by high neighbors ('Pocket of Poverty')."
], CRIMSON)

# =============================================================================
# SLIDE 9: LISA WEALTH CLUSTERS MAP (WITH IMAGE)
# =============================================================================
s9 = prs.slides.add_slide(blank_layout)
add_header(s9, "National Wealth Clusters: Anselin LISA Map", "EMPIRICAL FINDINGS")
add_image_safe(s9, 'docs/figures/05_moran_and_lisa_clusters.png', 0.8, 1.4, 11.7, 5.6)

# =============================================================================
# SLIDE 10: SECTOR 1 - HEALTHCARE DESERTS (QUESTION -> ANSWER)
# =============================================================================
s10 = prs.slides.add_slide(blank_layout)
add_header(s10, "Sector 1: Public Health & Healthcare Deserts", "SECTOR ANALYSIS")
add_text_box(s10, 0.8, 1.4, 4.5, 5.6, "The Question & Method", [
    "THE QUESTION:",
    "Where are populations trapped with zero primary healthcare facilities?",
    "THE DATA EXPLORATION:",
    "We cross-referenced gridded population against 35,000+ registered health facilities across 9,308 wards.",
    "THE SPATIAL STATS ANSWER:",
    "We identified 1,800+ Healthcare Deserts: wards with over 17,000 people and 0 registered clinics.",
    "Over 30 million Nigerians live in these desert wards."
], CRIMSON)
add_image_safe(s10, 'docs/figures/01_national_wealth_and_health_deserts.png', 5.5, 1.4, 7.0, 5.6)

# =============================================================================
# SLIDE 11: SECTOR 2 - DISEASE EPIDEMIOLOGY: MALARIA (QUESTION -> ANSWER)
# =============================================================================
s11 = prs.slides.add_slide(blank_layout)
add_header(s11, "Sector 2: Disease Epidemiology & Malaria Surveillance", "SECTOR ANALYSIS")
add_text_box(s11, 0.8, 1.4, 4.5, 5.6, "The Question & Method", [
    "THE QUESTION:",
    "Does malaria transmission cluster in distinct environmental corridors, and do local clinics buffer against the disease?",
    "THE SPATIAL STATS ANSWER:",
    "Malaria prevalence exhibits extreme clustering (Moran's I = 0.886, p = 0.001).",
    "Southern mangrove and riverine basins form massive endemic hotspots (PfPR > 40%).",
    "Wards with higher health clinic rates experience significantly lower active parasite rates."
], NAVY)
add_image_safe(s11, 'docs/figures/07_disease_epidemiology_malaria_map.png', 5.5, 1.4, 7.0, 5.6)

# =============================================================================
# SLIDE 12: SECTOR 3 - GEOMARKETING & COMMERCIAL RETAIL (QUESTION -> ANSWER)
# =============================================================================
s12 = prs.slides.add_slide(blank_layout)
add_header(s12, "Sector 3: Geomarketing & Retail Expansion Strategy", "SECTOR ANALYSIS")
add_text_box(s12, 0.8, 1.4, 4.5, 5.6, "The Question & Method", [
    "THE QUESTION:",
    "Where can modern supermarkets, retail chains, and fintechs expand without facing crowded competition?",
    "THE SPATIAL STATS ANSWER:",
    "We cross-referenced the Relative Wealth Index against physical market density.",
    "Tier 2 Wards (High Wealth, Low Markets) represent prime expansion targets.",
    "These wards have high consumer purchasing power but an absence of formal retail competition."
], CORAL)
add_image_safe(s12, 'docs/figures/02_commercial_retail_strategy.png', 5.5, 1.4, 7.0, 5.6)

# =============================================================================
# SLIDE 13: SECTOR 4 - CULTURAL GEOGRAPHY & PEACEBUILDING (QUESTION -> ANSWER)
# =============================================================================
s13 = prs.slides.add_slide(blank_layout)
add_header(s13, "Sector 4: Cultural Geography & Community Cohesion", "SECTOR ANALYSIS")
add_text_box(s13, 0.8, 1.4, 4.5, 5.6, "The Question & Method", [
    "THE QUESTION:",
    "How are faith institutions distributed, and where are the cultural transition zones of high co-presence?",
    "THE SPATIAL STATS ANSWER:",
    "Churches and mosques sort into a sharp North-South divide (Moran's I = 0.812).",
    "The Shannon Entropy Index isolates the Middle Belt (Plateau, Benue, Kaduna South) as intense zones of religious coexistence.",
    "Peacebuilding and vaccination drives must engage dominant faith leaders in each ward."
], TEAL)
add_image_safe(s13, 'docs/figures/03_religious_cultural_geography.png', 5.5, 1.4, 7.0, 5.6)

# =============================================================================
# SLIDE 14: SECTOR 5 - WASH & CLEAN WATER EQUITY (QUESTION -> ANSWER)
# =============================================================================
s14 = prs.slides.add_slide(blank_layout)
add_header(s14, "Sector 5: Public Utilities (WASH) & Infrastructure Equity", "SECTOR ANALYSIS")
add_text_box(s14, 0.8, 1.4, 4.5, 5.6, "The Question & Method", [
    "THE QUESTION:",
    "How unequal is the distribution of clean water points across Nigeria, and which communities need urgent boreholes?",
    "THE SPATIAL STATS ANSWER:",
    "Lorenz Curve analysis reveals severe inequality: Gini = 0.72 for clean water points.",
    "The Ward Priority Index (WPI) synthesizes water deficit, health deficit, and poverty into 4 Action Tiers.",
    "Directs water boards to drill solar boreholes in Tier 1 Emergency wards first."
], NAVY)
add_image_safe(s14, 'docs/figures/06_infrastructure_inequality_and_ward_priority_tiers.png', 5.5, 1.4, 7.0, 5.6)

# =============================================================================
# SLIDE 15: SPATIAL ECONOMETRICS - WHY OLS FAILS
# =============================================================================
s15 = prs.slides.add_slide(blank_layout)
add_header(s15, "Spatial Econometrics: Why Ordinary Least Squares (OLS) Fails", "ECONOMETRIC THEORY")
add_text_box(s15, 0.8, 1.4, 5.7, 5.6, "The Classical OLS Regression", [
    "Standard Equation:",
    "   y = X*beta + e",
    "OLS assumes errors e are independent:",
    "   Cov(e_i, e_j) = 0 for i != j",
    "When geographic units are tested, this assumption breaks down.",
    "Errors cluster spatially: poor wards border poor wards, and wealthy wards border wealthy wards."
], CRIMSON)
add_text_box(s15, 6.8, 1.4, 5.7, 5.6, "The Consequences in Real Life", [
    "If Spatial Lag is omitted:",
    "Your beta estimates are BIASED and WRONG.",
    "If Spatial Error is present:",
    "Your standard errors are too small, making you believe a policy or marketing variable works when it actually doesn't (Type-I Error).",
    "We must test OLS residuals using Anselin's Lagrange Multiplier (LM) diagnostics."
], NAVY)

# =============================================================================
# SLIDE 16: SPATIAL LAG (SAR) & THE SPATIAL MULTIPLIER
# =============================================================================
s16 = prs.slides.add_slide(blank_layout)
add_header(s16, "The Spatial Lag Model (SAR) & The 2.4x Multiplier", "ECONOMETRIC THEORY")
add_text_box(s16, 0.8, 1.4, 5.7, 5.6, "The SAR Mathematical Model", [
    "Equation:",
    "   y = rho * W*y + X*beta + e",
    "rho is the spatial autoregressive coefficient.",
    "W*y is the spatial lag (neighboring wealth/health).",
    "Reduced Form Equation:",
    "   y = (I - rho*W)^(-1) * X*beta + (I - rho*W)^(-1) * e",
    "The matrix (I - rho*W)^(-1) is the Spatial Multiplier."
], NAVY)
add_text_box(s16, 6.8, 1.4, 5.7, 5.6, "The 2.4x Multiplier Explained", [
    "Our empirical regression on 9,308 wards found:",
    "   rho = 0.5842  (p < 0.0001)",
    "Spatial Multiplier:",
    "   1 / (1 - rho) = 1 / (1 - 0.5842) = 2.405x",
    "In simple words:",
    "For every 1.0 unit of wealth created in a focal ward by building a market or clinic, an additional 1.405 units spill over into surrounding wards!",
    "This justifies regional joint funding between neighboring LGAs."
], CORAL)

# =============================================================================
# SLIDE 17: SPATIAL ERROR MODEL (SEM)
# =============================================================================
s17 = prs.slides.add_slide(blank_layout)
add_header(s17, "The Spatial Error Model (SEM): Unobserved Geographic Shocks", "ECONOMETRIC THEORY")
add_text_box(s17, 0.8, 1.4, 5.7, 5.6, "The SEM Mathematical Model", [
    "Equation:",
    "   y = X*beta + u",
    "where the disturbance u is spatially autocorrelated:",
    "   u = lambda * W*u + e",
    "lambda is the spatial error coefficient.",
    "e is classical independent white noise."
], NAVY)
add_text_box(s17, 6.8, 1.4, 5.7, 5.6, "When to Use SEM", [
    "Use SEM when spatial dependence is not a behavioral spillover, but caused by unobserved geographic factors:",
    "• Regional climate and rainfall variations",
    "• Soil fertility and topography",
    "• Historical state-level policies",
    "Results for Nigeria: lambda = 0.6124 (p < 0.0001). Correcting this error improves model fit by thousands of AIC points."
], TEAL)

# =============================================================================
# SLIDE 18: ECONOMETRIC MODEL COMPARISON TABLE
# =============================================================================
s18 = prs.slides.add_slide(blank_layout)
add_header(s18, "Comparing Econometric Models Across 9,308 Wards", "EMPIRICAL RESULTS")
add_text_box(s18, 0.8, 1.4, 11.7, 5.6, "Model Performance Summary", [
    "MODEL 1: Classical OLS",
    "   Pseudo R2 = 0.2841 | Log-Likelihood = -5,812.4 | AIC = 11,636.8",
    "   Moran's I on Residuals = 0.421 (p < 0.001) -> Severe residual spatial autocorrelation.",
    "",
    "MODEL 2: Spatial Lag Model (SAR)",
    "   Pseudo R2 = 0.5318 | Log-Likelihood = -3,941.2 | AIC = 7,896.4",
    "   rho = 0.5842 (p < 0.0001) -> Huge improvement; AIC drops by 3,740 points!",
    "",
    "MODEL 3: Spatial Error Model (SEM)",
    "   Pseudo R2 = 0.5462 | Log-Likelihood = -3,884.6 | AIC = 7,781.2",
    "   lambda = 0.6124 (p < 0.0001) -> Optimal specification for structural unobserved spatial shocks.",
    "",
    "TAKEAWAY: Spatial models explain almost twice as much variance as classical OLS."
], NAVY)

# =============================================================================
# SLIDE 19: 15 USE CASES TAXONOMY (PART 1)
# =============================================================================
s19 = prs.slides.add_slide(blank_layout)
add_header(s19, "15 Real-World Enterprise & Policy Use Cases (Part 1)", "TAXONOMY")
add_text_box(s19, 0.8, 1.4, 5.7, 5.6, "Public Health & Human Capital", [
    "1. Disease Surveillance: Isolating malaria and cholera transmission clusters for bed net distribution.",
    "2. Healthcare Facility Siting: Eliminating 1,800+ Healthcare Deserts to protect mothers and infants.",
    "3. Clean Water (WASH): Targeting public boreholes using Lorenz spatial inequality curves.",
    "4. Education Planning: Mapping school dropout risk zones and child density catchments."
], NAVY)
add_text_box(s19, 6.8, 1.4, 5.7, 5.6, "Commercial Retail & Fintech", [
    "5. Commercial Supermarkets: Siting stores in affluent, under-retailed wards (Tier 2).",
    "6. Fintech POS Agents: Placing cash-in/cash-out agency banking kiosks in unbanked cash corridors.",
    "7. Telecoms 4G/5G Towers: Balancing population density and signal distance decay.",
    "8. Off-Grid Solar Mini-Grids: Identifying affluent un-electrified settlements for private solar power."
], TEAL)

# =============================================================================
# SLIDE 20: 15 USE CASES TAXONOMY (PART 2)
# =============================================================================
s20 = prs.slides.add_slide(blank_layout)
add_header(s20, "15 Real-World Enterprise & Policy Use Cases (Part 2)", "TAXONOMY")
add_text_box(s20, 0.8, 1.4, 5.7, 5.6, "Civic, Safety & Environment", [
    "9. Cultural Peacebuilding: Using Shannon Entropy to target inter-religious dialogue in the Middle Belt.",
    "10. Emergency Response: Optimizing fire and police station transit radiuses.",
    "11. Climate Heat Islands: Detecting urban heat stress vs. tree canopy deficits.",
    "12. Agricultural Supply Chains: Locating wholesale grain silos along farm-to-market corridors."
], CORAL)
add_text_box(s20, 6.8, 1.4, 5.7, 5.6, "Urban Infrastructure & Governance", [
    "13. Real Estate Valuation: Spatial hedonic modeling of neighborhood infrastructure premiums.",
    "14. Transport Corridors: Measuring spatial economic multipliers (2.4x) for road projects.",
    "15. Electoral Logistics: Routing polling units to minimize voter travel distance and lines."
], NAVY)

# =============================================================================
# SLIDE 21: CONCLUSION & ACTION ROADMAP
# =============================================================================
s21 = prs.slides.add_slide(blank_layout)
add_header(s21, "Strategic Summary: How to Use This in Your Organization", "ACTION PLAN")
add_text_box(s21, 0.8, 1.4, 5.7, 5.6, "Key Rules for Decision-Makers", [
    "1. Stop Using Flat Averages: Look at ward-level spatial distributions.",
    "2. Account for Spillovers: Remember that the true return on capital is 2.4x due to neighboring gains.",
    "3. Start with a Question: Formulate a clear hypothesis before running any spatial model.",
    "4. Spatial Stats First: Use spatial lags and spatial CV before applying Machine Learning."
], NAVY)
add_text_box(s21, 6.8, 1.4, 5.7, 5.6, "How to Access & Run This Project", [
    "• Fast Track (5 seconds): Unpack data/processed_data_bundle.zip (9.2 MB) and open any notebook.",
    "• Masterclasses Available in notebooks/:",
    "   - 00: Master Executive Handbook & Inspector",
    "   - 01: Multi-Sector ESDA & Malaria Surveillance",
    "   - 02: Spatial Econometrics, SAR/SEM & Multipliers",
    "   - 03: Decision Intelligence Across 15 Use Cases",
    "• GitHub: https://github.com/SammyGIS/dsn_advanced_spatial_stats.git"
], TEAL)

prs_file = 'docs/spatial_statistics_masterclass_presentation.pptx'
prs.save(prs_file)
print(f"Successfully generated {prs_file} with {len(prs.slides)} slides and embedded map figures!")
